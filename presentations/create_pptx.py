#!/usr/bin/env python3
"""Tạo slide thuyết trình Chương 1 – Tổ chức sự kiện."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import nsmap
from pptx.oxml import parse_xml
from pptx.util import Emu, Inches, Pt

ASSETS = Path(__file__).resolve().parent / "assets"
OUT = Path(__file__).resolve().parent / "Chuong1_To_Chuc_Su_Kien.pptx"

# Palette: ocean teal academic (phù hợp lễ hội biển / DIFF)
TEAL = RGBColor(0x0B, 0x4F, 0x6C)
TEAL_DARK = RGBColor(0x07, 0x33, 0x45)
TEAL_SOFT = RGBColor(0x1B, 0x7A, 0x9A)
ACCENT = RGBColor(0xD4, 0x8A, 0x2C)  # amber gold
CREAM = RGBColor(0xF7, 0xF5, 0xF0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1A, 0x23, 0x32)
MUTED = RGBColor(0x4A, 0x55, 0x63)
LIGHT_TEAL = RGBColor(0xE6, 0xF2, 0xF5)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
SOFT_LINE = RGBColor(0xC5, 0xD6, 0xDC)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def set_run(run, size=18, bold=False, color=INK, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_textbox(slide, left, top, width, height, text, size=18, bold=False,
                color=INK, align=PP_ALIGN.LEFT, font="Calibri", anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color, font=font)
    return box


def add_multiline(slide, left, top, width, height, lines, size=15, color=INK,
                  bold_first=False, bullet=False, spacing=8):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(spacing)
        run = p.add_run()
        prefix = "• " if bullet else ""
        run.text = prefix + line
        set_run(run, size=size, bold=(bold_first and i == 0), color=color)
    return box


def rect(slide, left, top, width, height, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
    return shape


def rounded(slide, left, top, width, height, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    # softer corners
    try:
        shape.adjustments[0] = 0.12
    except Exception:
        pass
    return shape


def accent_bar(slide, left=Inches(0.5), top=Inches(1.15), width=Inches(1.2)):
    return rect(slide, left, top, width, Inches(0.08), ACCENT)


def section_header(slide, section, title, subtitle=None):
    rect(slide, 0, 0, SLIDE_W, Inches(0.18), TEAL)
    add_textbox(slide, Inches(0.55), Inches(0.35), Inches(12), Inches(0.35),
                section, size=13, bold=True, color=TEAL_SOFT)
    add_textbox(slide, Inches(0.55), Inches(0.65), Inches(12), Inches(0.55),
                title, size=28, bold=True, color=TEAL_DARK)
    accent_bar(slide)
    if subtitle:
        add_textbox(slide, Inches(0.55), Inches(1.3), Inches(12), Inches(0.4),
                    subtitle, size=14, color=MUTED)


def footer(slide, page, total=16):
    add_textbox(slide, Inches(0.55), Inches(7.1), Inches(9), Inches(0.3),
                "Chương 1 · Tổ chức sự kiện · Mục 1.3–1.6", size=10, color=MUTED)
    add_textbox(slide, Inches(11.5), Inches(7.1), Inches(1.3), Inches(0.3),
                f"{page}/{total}", size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def card(slide, left, top, width, height, title, body_lines, title_size=14, body_size=12):
    rounded(slide, left, top, width, height, WHITE)
    # left accent strip
    rect(slide, left, top, Inches(0.08), height, TEAL)
    add_textbox(slide, left + Inches(0.22), top + Inches(0.15),
                width - Inches(0.35), Inches(0.35),
                title, size=title_size, bold=True, color=TEAL)
    add_multiline(slide, left + Inches(0.22), top + Inches(0.5),
                  width - Inches(0.35), height - Inches(0.6),
                  body_lines, size=body_size, color=INK, bullet=False, spacing=4)


def blank_slide(prs):
    layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(layout)
    rect(slide, 0, 0, SLIDE_W, SLIDE_H, CREAM)
    return slide


def add_picture_safe(slide, path, left, top, width=None, height=None):
    path = Path(path)
    if not path.exists():
        return None
    kwargs = {}
    if width is not None:
        kwargs["width"] = width
    if height is not None:
        kwargs["height"] = height
    return slide.shapes.add_picture(str(path), left, top, **kwargs)


# ---------- slides ----------

def slide_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, SLIDE_W, SLIDE_H, TEAL_DARK)
    add_picture_safe(slide, ASSETS / "cover-diff-festival.jpg",
                     Inches(5.2), 0, width=Inches(8.2), height=SLIDE_H)
    rect(slide, 0, 0, Inches(6.4), SLIDE_H, TEAL_DARK)
    rect(slide, Inches(6.2), 0, Inches(0.2), SLIDE_H, ACCENT)

    add_textbox(slide, Inches(0.7), Inches(1.5), Inches(5.2), Inches(0.4),
                "MÔN HỌC · TỔ CHỨC SỰ KIỆN", size=13, bold=True, color=ACCENT)
    box = slide.shapes.add_textbox(Inches(0.7), Inches(2.05), Inches(5.2), Inches(1.7))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(["Chương 1", "Cơ sở lý thuyết &", "phân tích thực tế"]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line
        set_run(run, size=30, bold=True, color=WHITE)
        p.space_after = Pt(2)

    add_multiline(
        slide, Inches(0.7), Inches(4.0), Inches(5.2), Inches(0.9),
        ["Mục 1.3 – 1.6", "Thành phần · Yếu tố ảnh hưởng · Tác động · Sự kiện online"],
        size=14, color=RGBColor(0xC8, 0xD8, 0xE0), spacing=4
    )
    add_multiline(
        slide, Inches(0.7), Inches(5.3), Inches(5.2), Inches(0.8),
        ["Ví dụ: DIFF 2026 · Festival Biển Khánh Hòa", "Festival Huế · WWDC26"],
        size=13, color=RGBColor(0xA8, 0xC0, 0xCC), spacing=4
    )
    add_textbox(slide, Inches(0.7), Inches(6.5), Inches(5.2), Inches(0.35),
                "Thuyết trình học thuật 2026", size=12, color=ACCENT)


def slide_toc(prs):
    slide = blank_slide(prs)
    section_header(slide, "NỘI DUNG THUYẾT TRÌNH", "Mục lục chương")
    items = [
        ("01", "1.3.1", "Các thành phần tham gia — DIFF 2026"),
        ("02", "1.3.2", "Thị trường tổ chức sự kiện tại Việt Nam"),
        ("03", "1.4.1", "Yếu tố vĩ mô — Festival Biển Khánh Hòa"),
        ("04", "1.4.2", "Yếu tố vi mô — Lễ tốt nghiệp (giả định)"),
        ("05", "1.5.1", "Tác động tích cực — Festival Huế 2026"),
        ("06", "1.5.2", "Tác động tiêu cực & hạn chế — DIFF"),
        ("07", "1.6", "Sự kiện online / hybrid — WWDC26"),
        ("08", "KL", "Kết luận chương"),
    ]
    for i, (num, code, title) in enumerate(items):
        col = i // 4
        row = i % 4
        left = Inches(0.55 + col * 6.3)
        top = Inches(1.7 + row * 1.15)
        rounded(slide, left, top, Inches(5.9), Inches(0.95), WHITE)
        rect(slide, left, top, Inches(0.9), Inches(0.95), TEAL)
        add_textbox(slide, left, top + Inches(0.25), Inches(0.9), Inches(0.45),
                    num, size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_textbox(slide, left + Inches(1.1), top + Inches(0.18), Inches(4.5), Inches(0.3),
                    code, size=12, bold=True, color=TEAL_SOFT)
        add_textbox(slide, left + Inches(1.1), top + Inches(0.45), Inches(4.5), Inches(0.4),
                    title, size=14, color=INK)
    footer(slide, 2)


def slide_stakeholders(prs):
    slide = blank_slide(prs)
    section_header(
        slide, "1.3.1 CÁC THÀNH PHẦN THAM GIA",
        "Ai tham gia tổ chức sự kiện?",
        "Ví dụ xuyên suốt: Lễ hội pháo hoa quốc tế Đà Nẵng (DIFF) 2026"
    )
    add_picture_safe(slide, ASSETS / "stakeholders-network.jpg",
                     Inches(8.3), Inches(1.7), width=Inches(4.5))
    parties = [
        ("Chủ sự kiện / phối hợp", "UBND Đà Nẵng + Sun Group: mục tiêu, nguồn lực, thẩm quyền"),
        ("Nhà tài trợ", "Tiền / sản phẩm / dịch vụ ↔ quyền lợi truyền thông"),
        ("Đơn vị tổ chức", "Kịch bản, lịch 6 đêm, điều phối & rủi ro"),
        ("Nhà cung ứng", "Sân khấu, âm thanh, an ninh, hậu cần"),
        ("Khán giả & truyền thông", "Trải nghiệm hành trình + khu tác nghiệp"),
        ("Chính quyền & cộng đồng", "Giao thông, trật tự, sinh hoạt đô thị"),
    ]
    for i, (t, b) in enumerate(parties):
        col = i % 2
        row = i // 2
        left = Inches(0.5 + col * 3.85)
        top = Inches(1.75 + row * 1.55)
        card(slide, left, top, Inches(3.7), Inches(1.4), t, [b], title_size=13, body_size=12)
    footer(slide, 3)


def slide_stakeholders_lesson(prs):
    slide = blank_slide(prs)
    section_header(
        slide, "1.3.1 PHÂN TÍCH SÂU — DIFF 2026",
        "Bài học quản trị các bên liên quan"
    )
    points = [
        ("Không đồng nhất vai trò",
         "Chủ sự kiện ≠ nhà tài trợ. Đơn vị tổ chức không nhất thiết tự làm mọi dịch vụ."),
        ("Thống nhất mục tiêu sớm",
         "Quảng bá du lịch phải khớp giới hạn hạ tầng & an toàn; thống nhất chỉ số đánh giá."),
        ("Điều hành đa đêm",
         "DIFF 6 đêm (30/5–11/7/2026): lịch diễn, tổng duyệt, kỹ thuật, truyền thông theo chuỗi."),
        ("Đo bằng hành trình",
         "Chất lượng = tiết mục + đón khách + di chuyển + chỉ dẫn + an toàn."),
        ("Ràng buộc nhà tài trợ",
         "Vị trí quảng bá, thời gian, quyền lợi & nghĩa vụ trong thỏa thuận."),
        ("Đầu mối chung",
         "Lập sơ đồ stakeholder, bảng phân công và kênh liên lạc thống nhất."),
    ]
    for i, (t, b) in enumerate(points):
        col = i % 3
        row = i // 3
        left = Inches(0.5 + col * 4.2)
        top = Inches(1.7 + row * 2.4)
        rounded(slide, left, top, Inches(4.0), Inches(2.15), WHITE)
        rect(slide, left, top, Inches(4.0), Inches(0.12), ACCENT if row == 0 else TEAL)
        add_textbox(slide, left + Inches(0.25), top + Inches(0.35), Inches(3.5), Inches(0.45),
                    t, size=15, bold=True, color=TEAL)
        add_textbox(slide, left + Inches(0.25), top + Inches(0.85), Inches(3.5), Inches(1.1),
                    b, size=13, color=INK)
    footer(slide, 4)


def slide_market(prs):
    slide = blank_slide(prs)
    section_header(
        slide, "1.3.2 THỊ TRƯỜNG TỔ CHỨC SỰ KIỆN",
        "Thị trường Việt Nam & vị trí của DIFF"
    )
    types = [
        "Hội nghị – hội thảo", "Du lịch MICE", "Hội chợ – triển lãm",
        "Lễ hội văn hóa – du lịch", "Thể thao", "Nghệ thuật",
        "Sự kiện thương hiệu", "Online / Hybrid",
    ]
    for i, t in enumerate(types):
        col = i % 4
        row = i // 4
        left = Inches(0.5 + col * 3.15)
        top = Inches(1.65 + row * 0.85)
        rounded(slide, left, top, Inches(3.0), Inches(0.7), WHITE)
        add_textbox(slide, left + Inches(0.15), top + Inches(0.18), Inches(2.7), Inches(0.4),
                    t, size=14, bold=True, color=TEAL, align=PP_ALIGN.CENTER)

    rounded(slide, Inches(0.5), Inches(3.6), Inches(12.3), Inches(3.0), WHITE)
    rect(slide, Inches(0.5), Inches(3.6), Inches(0.12), Inches(3.0), ACCENT)
    add_textbox(slide, Inches(0.9), Inches(3.8), Inches(11.5), Inches(0.4),
                "DIFF 2026 trong thị trường", size=16, bold=True, color=TEAL)
    add_multiline(
        slide, Inches(0.9), Inches(4.3), Inches(11.5), Inches(2.1),
        [
            "Giao điểm: lễ hội – biểu diễn nghệ thuật – du lịch.",
            "Chuỗi dịch vụ nhiều nhà cung cấp: trình diễn, kỹ thuật, quảng bá, bán vé, đón tiếp.",
            "Quy mô đa đêm → chuyên môn hóa cao; chi phí điều phối cũng tăng.",
            "Bài học: đánh giá năng lực nhà thầu, tiêu chuẩn dịch vụ, đầu mối liên hệ chung.",
            "Lưu ý: không dùng nhận định cũ “VN chưa có ngành tổ chức sự kiện độc lập” cho 2026.",
        ],
        size=14, bullet=True, spacing=6
    )
    footer(slide, 5)


def slide_macro_intro(prs):
    slide = blank_slide(prs)
    section_header(
        slide, "1.4.1 CÁC YẾU TỐ VĨ MÔ",
        "Sáu nhóm môi trường bên ngoài",
        "Ví dụ: Festival Biển Khánh Hòa 2026 (17–19/7) — “Sắc màu đại dương – Vươn tầm quốc tế”"
    )
    add_picture_safe(slide, ASSETS / "festival-bien-khanh-hoa.jpg",
                     Inches(7.8), Inches(1.7), width=Inches(5.0))
    factors = [
        ("01", "Nhân khẩu học"),
        ("02", "Kinh tế"),
        ("03", "Tự nhiên"),
        ("04", "Công nghệ"),
        ("05", "Chính trị – pháp luật"),
        ("06", "Văn hóa"),
    ]
    for i, (n, name) in enumerate(factors):
        top = Inches(1.7 + i * 0.8)
        rounded(slide, Inches(0.5), top, Inches(6.9), Inches(0.7), WHITE)
        rect(slide, Inches(0.5), top, Inches(0.85), Inches(0.7), TEAL)
        add_textbox(slide, Inches(0.5), top + Inches(0.15), Inches(0.85), Inches(0.4),
                    n, size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_textbox(slide, Inches(1.55), top + Inches(0.18), Inches(5.5), Inches(0.4),
                    name, size=16, bold=True, color=INK)
    footer(slide, 6)


def slide_macro_analysis(prs):
    slide = blank_slide(prs)
    section_header(
        slide, "1.4.1 PHÂN TÍCH — FESTIVAL BIỂN KHÁNH HÒA",
        "Tác động qua lại của 6 yếu tố vĩ mô"
    )
    data = [
        ("Nhân khẩu học", "Carnival & văn hóa thu hút nhiều nhóm tuổi → phân nhóm khán giả, chỉ dẫn, khu hỗ trợ."),
        ("Kinh tế", "AR Realtime, 3D Mapping, truyền hình → ngân sách theo hạng mục + dự phòng thử nghiệm."),
        ("Tự nhiên", "Ven biển đẹp nhưng rủi ro mưa gió → ngưỡng dừng/hoãn, trú tránh, bảo vệ thiết bị."),
        ("Công nghệ", "Hiệu ứng thị giác mạnh; cần đồng bộ tín hiệu, kiểm thử nhiều vòng, kỹ thuật trực."),
        ("Chính trị – PL", "Diễu hành ngoài trời: an toàn, bản quyền, tiếng ồn → danh mục pháp lý sớm."),
        ("Văn hóa", "Cân bằng kỹ xảo hiện đại với bản sắc biển Khánh Hòa — tránh thông điệp nông."),
    ]
    for i, (t, b) in enumerate(data):
        col = i % 3
        row = i // 3
        left = Inches(0.45 + col * 4.25)
        top = Inches(1.65 + row * 2.5)
        card(slide, left, top, Inches(4.1), Inches(2.25), t, [b], title_size=15, body_size=13)
    footer(slide, 7)


def slide_micro(prs):
    slide = blank_slide(prs)
    section_header(
        slide, "1.4.2 CÁC YẾU TỐ VI MÔ",
        "Nhân sự · Ngân sách · Nhà cung ứng · Địa điểm",
        "Tình huống giả định: Lễ tốt nghiệp 500 sinh viên & người thân"
    )
    add_picture_safe(slide, ASSETS / "graduation-event.jpg",
                     Inches(8.4), Inches(1.7), width=Inches(4.4))
    items = [
        ("Nhân sự", "2 bàn đón tiếp dễ ùn → chia làn, check-in QR, trưởng ca."),
        ("Ngân sách", "Tách cố định / biến đổi / dự phòng; người duyệt phát sinh."),
        ("Nhà cung ứng", "Màn hình xa hỏng → chạy thử + thiết bị dự phòng."),
        ("Địa điểm", "Sảnh nhỏ = điểm nghẽn; khảo sát sức chứa & lối đi."),
        ("Người tham dự", "Tập trung cùng giờ → khung giờ / chỉ dẫn rõ."),
        ("Phối hợp ngoài", "Bảo vệ – y tế – kỹ thuật: một đầu mối điều hành."),
    ]
    for i, (t, b) in enumerate(items):
        top = Inches(1.65 + i * 0.8)
        rounded(slide, Inches(0.5), top, Inches(7.6), Inches(0.72), WHITE)
        add_textbox(slide, Inches(0.7), top + Inches(0.08), Inches(2.2), Inches(0.5),
                    t, size=14, bold=True, color=TEAL)
        add_textbox(slide, Inches(2.9), top + Inches(0.15), Inches(5.0), Inches(0.5),
                    b, size=13, color=INK)
    footer(slide, 8)


def slide_positive(prs):
    slide = blank_slide(prs)
    section_header(
        slide, "1.5.1 TÁC ĐỘNG TÍCH CỰC",
        "Giá trị kinh tế · văn hóa · xã hội · điểm đến",
        "Ví dụ: Festival Huế 2026 — hoạt động theo bốn mùa"
    )
    add_picture_safe(slide, ASSETS / "festival-hue.jpg",
                     Inches(0.5), Inches(1.7), width=Inches(5.5))
    impacts = [
        ("Kinh tế", "Kéo dài lưu trú → lưu trú, ẩm thực, di chuyển, mua sắm (cần dữ liệu hậu sự kiện)."),
        ("Văn hóa", "Nghệ thuật đường phố đưa di sản gần người xem; tôn trọng không gian lịch sử."),
        ("Xã hội", "Cơ hội tham gia cộng đồng nếu dễ tiếp cận & ít gián đoạn sinh hoạt."),
        ("Hình ảnh điểm đến", "Chuỗi bốn mùa tạo điểm nhấn xuyên năm — đo bằng hài lòng & khả năng quay lại."),
    ]
    for i, (t, b) in enumerate(impacts):
        top = Inches(1.7 + i * 1.2)
        card(slide, Inches(6.3), top, Inches(6.5), Inches(1.05), t, [b],
             title_size=14, body_size=12)
    footer(slide, 9)


def slide_negative(prs):
    slide = blank_slide(prs)
    section_header(
        slide, "1.5.2 TÁC ĐỘNG TIÊU CỰC & HẠN CHẾ",
        "Nguy cơ cần quản trị tại DIFF 2026",
        "Các tình huống dưới đây là nguy cơ phân tích — không phải sự cố đã xác nhận"
    )
    risks = [
        ("Ùn tắc", "Phân luồng, công bố điểm đến/đi, hỗ trợ tại nút giao.\nTheo dõi: thời gian chờ, mật độ."),
        ("Quá tải khán giả", "Kiểm soát sức chứa, lối ra/vào, phân tán khu xem.\nTheo dõi: mật độ từng khu."),
        ("Rác thải", "Thùng phân loại, đội thu gom sau từng đêm.\nTheo dõi: khối lượng & thời gian sạch."),
        ("Tiếng ồn / cư dân", "Thông báo sớm, kiểm soát âm thanh, kênh phản ánh."),
        ("Khẩn cấp", "Điểm y tế, lối tiếp cận, đầu mối chỉ huy, diễn tập."),
        ("Cân bằng lợi ích", "Kinh tế – văn hóa ↔ chi phí môi trường, hạ tầng, cộng đồng."),
    ]
    for i, (t, b) in enumerate(risks):
        col = i % 3
        row = i // 3
        left = Inches(0.45 + col * 4.25)
        top = Inches(1.7 + row * 2.45)
        rounded(slide, left, top, Inches(4.1), Inches(2.25), WHITE)
        rect(slide, left, top, Inches(4.1), Inches(0.5), TEAL if i < 5 else ACCENT)
        add_textbox(slide, left + Inches(0.2), top + Inches(0.1), Inches(3.7), Inches(0.35),
                    t, size=14, bold=True, color=WHITE)
        add_textbox(slide, left + Inches(0.2), top + Inches(0.7), Inches(3.7), Inches(1.4),
                    b, size=13, color=INK)
    footer(slide, 10)


def slide_online_concept(prs):
    slide = blank_slide(prs)
    section_header(
        slide, "1.6.1 SỰ KIỆN ONLINE",
        "Khái niệm · đặc điểm · mô hình hybrid",
        "Ví dụ: Apple WWDC26 (8–12/6/2026) — Keynote online + hoạt động tại Apple Park"
    )
    add_picture_safe(slide, ASSETS / "hybrid-online-event.jpg",
                     Inches(7.9), Inches(1.7), width=Inches(4.9))
    add_multiline(
        slide, Inches(0.55), Inches(1.7), Inches(7.0), Inches(2.0),
        [
            "Online ≠ chỉ livestream: live, video sẵn, Q&A, tương tác, thư viện nội dung.",
            "Hybrid = trải nghiệm tại chỗ + trải nghiệm số — thiết kế hành trình riêng cho từng nhóm.",
            "WWDC26: Keynote đa nền tảng + 100+ session video + ngày trực tiếp tại Apple Park.",
        ],
        size=14, bullet=True, spacing=10
    )
    cards = [
        ("Tiếp cận", "Người ở xa xem được; xem lại giảm lệch múi giờ."),
        ("Phân tầng nội dung", "Keynote tổng quát + phiên chuyên môn theo nhu cầu."),
        ("Đo lường gợi ý", "Lượt xem, thời lượng, phiên sâu, phản hồi."),
    ]
    for i, (t, b) in enumerate(cards):
        left = Inches(0.55 + i * 2.4)
        card(slide, left, Inches(4.4), Inches(2.25), Inches(2.1), t, [b],
             title_size=13, body_size=12)
    footer(slide, 11)


def slide_pros_cons(prs):
    slide = blank_slide(prs)
    section_header(
        slide, "1.6.2 ƯU & NHƯỢC ĐIỂM",
        "Nhìn qua mô hình WWDC26"
    )
    # two columns
    rounded(slide, Inches(0.5), Inches(1.65), Inches(6.0), Inches(5.0), WHITE)
    rect(slide, Inches(0.5), Inches(1.65), Inches(6.0), Inches(0.6), TEAL)
    add_textbox(slide, Inches(0.7), Inches(1.75), Inches(5.5), Inches(0.4),
                "Ưu điểm", size=18, bold=True, color=WHITE)
    add_multiline(
        slide, Inches(0.8), Inches(2.5), Inches(5.4), Inches(3.8),
        [
            "Tiếp cận toàn cầu, không cần di chuyển",
            "Giảm chi phí đi lại – lưu trú cho người tham dự",
            "Nội dung chuyên môn có thể xem lại",
            "Đo lường đăng ký, tương tác, thời lượng xem",
            "Phân tầng session giúp học theo nhu cầu",
        ],
        size=15, bullet=True, spacing=12
    )

    rounded(slide, Inches(6.8), Inches(1.65), Inches(6.0), Inches(5.0), WHITE)
    rect(slide, Inches(6.8), Inches(1.65), Inches(6.0), Inches(0.6), ACCENT)
    add_textbox(slide, Inches(7.0), Inches(1.75), Inches(5.5), Inches(0.4),
                "Nhược điểm", size=18, bold=True, color=WHITE)
    add_multiline(
        slide, Inches(7.1), Inches(2.5), Inches(5.4), Inches(3.8),
        [
            "Thiếu tương tác trực diện / networking tại chỗ",
            "Phụ thuộc đường truyền, nền tảng, phụ đề",
            "Chi phí sản xuất – vận hành kỹ thuật vẫn cao",
            "Dễ phân tán chú ý khi nội dung dài",
            "Bảo mật dữ liệu & bản quyền cần dự phòng",
        ],
        size=15, bullet=True, spacing=12
    )
    footer(slide, 12)


def slide_online_types(prs):
    slide = blank_slide(prs)
    section_header(
        slide, "1.6.3 CÁC LOẠI HÌNH ONLINE",
        "Webinar · Video conference · Trade show · Concert · Hybrid"
    )
    add_picture_safe(slide, ASSETS / "online-event-types.jpg",
                     Inches(0.5), Inches(1.6), width=Inches(5.8))
    types = [
        ("Webinar", "1–vài diễn giả → nhiều người; trọng tâm chia sẻ + Q&A."),
        ("Video conference", "Nhóm nhỏ thảo luận 2 chiều, đồng biên tập, ra quyết định."),
        ("Online trade show", "Gian hàng số: trưng bày, tài liệu, kết nối tư vấn."),
        ("Online concert", "Hình ảnh – âm thanh – cảm xúc; tương tác tùy thiết kế."),
        ("Hybrid event", "Mục tiêu riêng cho tại chỗ & từ xa — không chỉ 1 camera."),
    ]
    for i, (t, b) in enumerate(types):
        top = Inches(1.6 + i * 0.95)
        card(slide, Inches(6.6), top, Inches(6.2), Inches(0.85), t, [b],
             title_size=13, body_size=12)
    footer(slide, 13)


def slide_online_tips(prs):
    slide = blank_slide(prs)
    section_header(
        slide, "1.6.4 LƯU Ý TỔ CHỨC ONLINE HIỆU QUẢ",
        "Quy trình & KPI mẫu",
        "Tình huống giả định: Hội thảo hướng nghiệp cho sinh viên"
    )
    steps = [
        ("1. Mục tiêu SMART", "200 ĐK · ≥120 dự · ≥40% khảo sát · ≥4/5 hài lòng"),
        ("2. Nền tảng", "Sức chứa, Q&A, ghi hình, bảo mật, kiểm tra link"),
        ("3. Phân vai", "MC · Diễn giả · Kỹ thuật · Quản lý chat"),
        ("4. Kịch bản phút", "Mở đầu → trình bày → Q&A → khảo sát → tài liệu"),
        ("5. Chạy thử + B", "Mic/cam/share; mất mạng → Q&A sẵn / video dự phòng"),
        ("6. Đo lường", "Tỷ lệ dự = dự/ĐK; phản hồi = trả lời/dự × 100%"),
    ]
    for i, (t, b) in enumerate(steps):
        col = i % 3
        row = i // 3
        left = Inches(0.45 + col * 4.25)
        top = Inches(1.7 + row * 2.45)
        rounded(slide, left, top, Inches(4.1), Inches(2.2), WHITE)
        add_textbox(slide, left + Inches(0.25), top + Inches(0.35), Inches(3.6), Inches(0.5),
                    t, size=16, bold=True, color=TEAL)
        add_textbox(slide, left + Inches(0.25), top + Inches(1.0), Inches(3.6), Inches(0.9),
                    b, size=13, color=INK)
    footer(slide, 14)


def slide_conclusion(prs):
    slide = blank_slide(prs)
    section_header(slide, "KẾT LUẬN CHƯƠNG", "Ba thông điệp then chốt")
    msgs = [
        ("01", "Điều phối nhiều bên",
         "Tổ chức sự kiện = phối hợp stakeholder chịu tác động vĩ mô & vi mô."),
        ("02", "Lợi ích & chi phí",
         "Tạo giá trị kinh tế – văn hóa – xã hội nhưng gây áp lực hạ tầng & cộng đồng."),
        ("03", "Online / Hybrid",
         "Tiếp cận rộng phải kèm nội dung phù hợp, công nghệ, dự phòng và đo lường."),
    ]
    for i, (n, t, b) in enumerate(msgs):
        left = Inches(0.5 + i * 4.2)
        rounded(slide, left, Inches(1.8), Inches(4.0), Inches(3.6), WHITE)
        rect(slide, left, Inches(1.8), Inches(4.0), Inches(0.9), TEAL)
        add_textbox(slide, left, Inches(1.95), Inches(4.0), Inches(0.55),
                    n, size=22, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
        add_textbox(slide, left + Inches(0.25), Inches(2.95), Inches(3.5), Inches(0.6),
                    t, size=16, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
        add_textbox(slide, left + Inches(0.3), Inches(3.7), Inches(3.4), Inches(1.4),
                    b, size=14, color=INK, align=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(0.5), Inches(5.8), Inches(12.3), Inches(0.8),
                "Câu hỏi thảo luận: Với sự kiện quy mô đông người ngoài trời, ưu tiên chỉ số nào trước — "
                "trải nghiệm khán giả, an toàn cộng đồng, hay hiệu quả truyền thông điểm đến?",
                size=14, color=MUTED)
    footer(slide, 15)


def slide_thanks(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_picture_safe(slide, ASSETS / "cover-diff-festival.jpg",
                     0, 0, width=SLIDE_W, height=SLIDE_H)
    rect(slide, Inches(1.5), Inches(2.1), Inches(10.3), Inches(3.3), TEAL_DARK)
    rect(slide, Inches(1.5), Inches(2.1), Inches(10.3), Inches(0.12), ACCENT)
    add_textbox(slide, Inches(1.7), Inches(2.6), Inches(9.9), Inches(0.8),
                "Cảm ơn quý thầy cô và các bạn", size=34, bold=True,
                color=WHITE, align=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(1.7), Inches(3.5), Inches(9.9), Inches(0.45),
                "Chương 1 · Tổ chức sự kiện · Mục 1.3–1.6", size=16,
                color=ACCENT, align=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(1.7), Inches(4.2), Inches(9.9), Inches(0.45),
                "Sẵn sàng trao đổi & nhận câu hỏi", size=15,
                color=RGBColor(0xC8, 0xD8, 0xE0), align=PP_ALIGN.CENTER)


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    slide_cover(prs)
    slide_toc(prs)
    slide_stakeholders(prs)
    slide_stakeholders_lesson(prs)
    slide_market(prs)
    slide_macro_intro(prs)
    slide_macro_analysis(prs)
    slide_micro(prs)
    slide_positive(prs)
    slide_negative(prs)
    slide_online_concept(prs)
    slide_pros_cons(prs)
    slide_online_types(prs)
    slide_online_tips(prs)
    slide_conclusion(prs)
    slide_thanks(prs)

    prs.save(OUT)
    print(f"Saved: {OUT}")
    print(f"Slides: {len(prs.slides)}")


if __name__ == "__main__":
    main()
